import io
import re

from app.data import SKILL_KEYWORDS

SKILL_ALIASES = {
    "Python": ["python", "py", "python3"],
    "SQL": ["sql", "mysql", "ms sql", "sql server", "oracle sql", "sqlite"],
    "Power BI": ["power bi", "powerbi", "power-bi", "dax", "power query"],
    "Excel": ["excel", "ms excel", "microsoft excel", "xlookup", "vlookup", "pivot table"],
    "Statistics": ["statistics", "statistical analysis", "probability", "hypothesis testing"],
    "Machine Learning": ["machine learning", "ml", "classification", "regression", "clustering"],
    "Deep Learning": ["deep learning", "neural network", "cnn", "rnn", "lstm"],
    "Data Visualization": ["data visualization", "visualization", "matplotlib", "seaborn", "plotly"],
    "Pandas": ["pandas", "dataframe"],
    "NumPy": ["numpy", "np array"],
    "Scikit-learn": ["scikit-learn", "sklearn", "scikit learn"],
    "TensorFlow": ["tensorflow", "tf"],
    "PyTorch": ["pytorch", "torch"],
    "React": ["react", "reactjs", "react.js"],
    "JavaScript": ["javascript", "js", "ecmascript"],
    "TypeScript": ["typescript", "ts"],
    "HTML": ["html", "html5"],
    "CSS": ["css", "css3"],
    "Tailwind CSS": ["tailwind", "tailwind css"],
    "Node.js": ["node", "node.js", "nodejs", "express"],
    "FastAPI": ["fastapi", "fast api"],
    "Django": ["django"],
    "Flask": ["flask"],
    "PostgreSQL": ["postgresql", "postgres", "pgadmin"],
    "MongoDB": ["mongodb", "mongo db", "mongoose"],
    "Docker": ["docker", "containerization", "container"],
    "AWS": ["aws", "amazon web services", "ec2", "s3", "lambda"],
    "Git": ["git", "github", "gitlab", "version control"],
    "NLP": ["nlp", "natural language processing", "text mining"],
    "Computer Vision": ["computer vision", "opencv", "image processing"],
    "Cybersecurity": ["cybersecurity", "cyber security", "security testing"],
    "Linux": ["linux", "ubuntu", "shell scripting", "bash"],
    "Networking": ["networking", "tcp/ip", "dns", "http", "osi model"],
    "Java": ["java", "core java"],
    "C++": ["c++", "cpp"],
    "C#": ["c#", "c sharp"],
    "PHP": ["php", "laravel"],
    "Tableau": ["tableau"],
    "Looker": ["looker", "looker studio", "google data studio"],
    "Azure": ["azure", "microsoft azure"],
    "Google Cloud": ["gcp", "google cloud"],
    "Kubernetes": ["kubernetes", "k8s"],
    "Jenkins": ["jenkins", "ci/cd", "cicd"],
    "REST API": ["rest api", "restful api", "api development"],
    "Selenium": ["selenium", "automation testing"],
    "Figma": ["figma", "ui design"],
    "Firebase": ["firebase", "firestore"],
    "Redux": ["redux", "redux toolkit"],
    "Next.js": ["next.js", "nextjs", "next js"],
    "Express.js": ["express", "express.js", "expressjs"],
    "Spring Boot": ["spring boot", "springboot"],
}

PROJECT_HEADING_PATTERN = re.compile(
    r"(projects?|academic projects?|personal projects?|major projects?|mini projects?)",
    re.IGNORECASE,
)


def normalize_text(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[^a-z0-9+#./\s-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def extract_skills_from_text(resume_text: str) -> list[str]:
    normalized = normalize_text(resume_text)
    detected = set()

    for keyword, display_name in SKILL_KEYWORDS.items():
        aliases = SKILL_ALIASES.get(display_name, [keyword])
        if any(has_skill_alias(normalized, alias) for alias in aliases):
            detected.add(display_name)

    return sorted(detected)


def has_skill_alias(text: str, alias: str) -> bool:
    normalized_alias = re.escape(alias.lower()).replace(r"\ ", r"[\s./_-]+")
    return bool(re.search(rf"(?<![a-z0-9]){normalized_alias}(?![a-z0-9])", text))


def clean_extracted_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_projects_from_text(resume_text: str) -> list[str]:
    lines = [line.strip(" -•\t") for line in resume_text.splitlines() if line.strip()]
    projects = []
    in_project_section = False

    for line in lines:
        if PROJECT_HEADING_PATTERN.fullmatch(line.strip(": ").lower()):
            in_project_section = True
            continue

        if in_project_section and re.match(
            r"^(education|experience|skills|certifications|achievements|contact|summary)\b",
            line,
            re.IGNORECASE,
        ):
            in_project_section = False

        looks_like_project = bool(
            re.search(
                r"\b(project|system|dashboard|prediction|forecasting|analysis|classifier|detection|portal|app|application|website|model)\b",
                line,
                re.IGNORECASE,
            )
        )

        if (in_project_section or looks_like_project) and 5 <= len(line) <= 120:
            cleaned = re.sub(r"\s+", " ", line).strip(":- ")
            if cleaned and cleaned.lower() not in {item.lower() for item in projects}:
                projects.append(cleaned)

        if len(projects) >= 6:
            break

    return projects


def extract_text_from_txt(content: bytes) -> str:
    return content.decode("utf-8", errors="ignore")


def extract_text_from_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as error:
        raise RuntimeError("Install pypdf to read PDF resumes.") from error

    reader = PdfReader(io.BytesIO(content))
    text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in reader.pages)
    if not text.strip():
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return clean_extracted_text(text)


def extract_text_from_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ModuleNotFoundError as error:
        raise RuntimeError("Install python-docx to read Word resumes.") from error

    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return clean_extracted_text("\n".join(paragraphs + table_cells))


def extract_text_from_resume_file(filename: str, content: bytes) -> str:
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(content)
    if lower_name.endswith(".docx"):
        return extract_text_from_docx(content)
    if lower_name.endswith(".txt"):
        return clean_extracted_text(extract_text_from_txt(content))

    raise ValueError("Unsupported resume format. Upload PDF, DOCX, or TXT.")
